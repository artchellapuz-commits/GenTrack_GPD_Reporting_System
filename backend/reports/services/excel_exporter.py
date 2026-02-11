try:
    import pandas as pd
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils.dataframe import dataframe_to_rows
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False
    
from datetime import datetime
import os
from django.conf import settings


class ExcelExporter:
    """Service class for generating Excel reports"""
    
    def __init__(self, queryset, report_type):
        if not EXCEL_AVAILABLE:
            raise ImportError("pandas and openpyxl are required for Excel export")
        
        self.queryset = queryset
        self.report_type = report_type
    
    def generate(self):
        """Generate Excel file based on report type"""
        if self.report_type == 'daily':
            return self._generate_daily_report()
        elif self.report_type == 'monthly':
            return self._generate_monthly_report()
        elif self.report_type == 'consolidated':
            return self._generate_consolidated_report()
    
    def _generate_daily_report(self):
        """Generate daily report"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Daily Generation Report"
        
        # Add header
        self._add_header(ws, "NPC DAILY GENERATION REPORT")
        
        # Column headers
        headers = ['Date', 'Plant', 'Unit', 'Generation (kWh)', 'Operating Hours', 
                   'Availability Hours', 'Forced Outage', 'Scheduled Outage', 
                   'Capacity Factor (%)', 'Availability Factor (%)']
        
        ws.append(headers)
        self._style_header_row(ws, ws.max_row)
        
        # Add data
        for report in self.queryset:
            ws.append([
                report.report_date.strftime('%Y-%m-%d'),
                report.plant.code,
                report.unit.unit_number,
                float(report.generation_kwh),
                float(report.operating_hours),
                float(report.availability_hours),
                float(report.forced_outage_hours),
                float(report.scheduled_outage_hours),
                float(report.capacity_factor) if report.capacity_factor else 0,
                float(report.availability_factor) if report.availability_factor else 0
            ])
        
        # Add totals
        self._add_totals(ws)
        
        # Auto-adjust column widths
        self._adjust_column_widths(ws)
        
        # Save file
        file_path = self._get_file_path('daily')
        wb.save(file_path)
        return file_path

    def _generate_monthly_report(self):
        """Generate monthly summary report"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Monthly Generation Report"
        
        # Add header
        self._add_header(ws, "NPC MONTHLY GENERATION REPORT")
        
        # Convert to DataFrame for easier aggregation
        data = []
        for report in self.queryset:
            data.append({
                'year': report.report_date.year,
                'month': report.report_date.month,
                'plant': report.plant.code,
                'unit': report.unit.unit_number,
                'generation_kwh': float(report.generation_kwh),
                'operating_hours': float(report.operating_hours),
                'forced_outage_hours': float(report.forced_outage_hours)
            })
        
        df = pd.DataFrame(data)
        
        # Group by year, month, plant, unit
        monthly = df.groupby(['year', 'month', 'plant', 'unit']).agg({
            'generation_kwh': 'sum',
            'operating_hours': 'sum',
            'forced_outage_hours': 'sum'
        }).reset_index()
        
        # Column headers
        headers = ['Year', 'Month', 'Plant', 'Unit', 'Total Generation (kWh)', 
                   'Total Operating Hours', 'Total Forced Outage Hours']
        ws.append(headers)
        self._style_header_row(ws, ws.max_row)
        
        # Add data
        for _, row in monthly.iterrows():
            ws.append([
                int(row['year']),
                int(row['month']),
                row['plant'],
                int(row['unit']),
                row['generation_kwh'],
                row['operating_hours'],
                row['forced_outage_hours']
            ])
        
        self._adjust_column_widths(ws)
        
        file_path = self._get_file_path('monthly')
        wb.save(file_path)
        return file_path
    
    def _generate_consolidated_report(self):
        """Generate consolidated report for all plants"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Consolidated Report"
        
        self._add_header(ws, "NPC CONSOLIDATED GENERATION REPORT - AGUS PLANTS")
        
        # Convert to DataFrame
        data = []
        for report in self.queryset:
            data.append({
                'date': report.report_date,
                'plant': report.plant.code,
                'generation_kwh': float(report.generation_kwh),
                'operating_hours': float(report.operating_hours)
            })
        
        df = pd.DataFrame(data)
        
        # Group by date and plant
        consolidated = df.groupby(['date', 'plant']).agg({
            'generation_kwh': 'sum',
            'operating_hours': 'sum'
        }).reset_index()
        
        # Pivot to show plants as columns
        pivot = consolidated.pivot(index='date', columns='plant', values='generation_kwh')
        pivot['Total'] = pivot.sum(axis=1)
        pivot = pivot.reset_index()
        
        # Add to worksheet
        for r in dataframe_to_rows(pivot, index=False, header=True):
            ws.append(r)
        
        self._style_header_row(ws, 1)
        self._adjust_column_widths(ws)
        
        file_path = self._get_file_path('consolidated')
        wb.save(file_path)
        return file_path
    
    def _add_header(self, ws, title):
        """Add formatted header to worksheet"""
        ws.append([title])
        ws.append([f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])
        ws.append([])
        
        # Style title
        ws['A1'].font = Font(size=14, bold=True)
        ws['A2'].font = Font(size=10, italic=True)
    
    def _style_header_row(self, ws, row_num):
        """Apply styling to header row"""
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        
        for cell in ws[row_num]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
    
    def _add_totals(self, ws):
        """Add totals row"""
        last_row = ws.max_row
        ws.append(['', '', 'TOTAL', 
                   f'=SUM(D4:D{last_row})',
                   f'=SUM(E4:E{last_row})',
                   f'=SUM(F4:F{last_row})',
                   f'=SUM(G4:G{last_row})',
                   f'=SUM(H4:H{last_row})',
                   f'=AVERAGE(I4:I{last_row})',
                   f'=AVERAGE(J4:J{last_row})'])
        
        # Style totals row
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    
    def _adjust_column_widths(self, ws):
        """Auto-adjust column widths"""
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    def _get_file_path(self, report_type):
        """Generate file path for export"""
        export_dir = os.path.join(settings.MEDIA_ROOT, 'exports')
        os.makedirs(export_dir, exist_ok=True)
        
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f"NPC_{report_type}_{timestamp}.xlsx"
        
        return os.path.join(export_dir, filename)
