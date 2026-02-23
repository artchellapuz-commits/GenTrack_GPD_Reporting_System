"""
Detailed PSR Template Analysis
Extracts EXACT formatting details from PSR REPORT-8AM.xlsx
"""
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Color

def rgb_to_hex(rgb):
    """Convert RGB to hex color"""
    if rgb and hasattr(rgb, 'rgb'):
        return rgb.rgb
    return None

def analyze_cell_format(cell):
    """Get detailed cell formatting"""
    info = {
        'value': cell.value,
        'font_name': cell.font.name if cell.font else None,
        'font_size': cell.font.size if cell.font else None,
        'font_bold': cell.font.bold if cell.font else False,
        'font_italic': cell.font.italic if cell.font else False,
        'font_color': rgb_to_hex(cell.font.color) if cell.font and cell.font.color else None,
        'fill_color': rgb_to_hex(cell.fill.fgColor) if cell.fill and cell.fill.fgColor else None,
        'alignment_horizontal': cell.alignment.horizontal if cell.alignment else None,
        'alignment_vertical': cell.alignment.vertical if cell.alignment else None,
        'alignment_wrap': cell.alignment.wrap_text if cell.alignment else False,
        'border_top': bool(cell.border.top.style) if cell.border and cell.border.top else False,
        'border_bottom': bool(cell.border.bottom.style) if cell.border and cell.border.bottom else False,
        'border_left': bool(cell.border.left.style) if cell.border and cell.border.left else False,
        'border_right': bool(cell.border.right.style) if cell.border and cell.border.right else False,
    }
    return info

# Load workbook
file_path = r'C:\Users\eladiong\Desktop\OJT Intern\REPORTS\8. PSR\PSR REPORT-8AM.xlsx'
print("=" * 80)
print("DETAILED PSR TEMPLATE ANALYSIS")
print("=" * 80)
print()

wb = openpyxl.load_workbook(file_path)
ws = wb.active

print(f"Sheet: {ws.title}")
print(f"Dimensions: {ws.dimensions}")
print()

# Analyze first 60 rows in detail
print("=" * 80)
print("DETAILED CELL ANALYSIS (Rows 1-60)")
print("=" * 80)
print()

for row_num in range(1, 61):
    row_has_content = False
    row_info = []
    
    for col_num in range(1, 15):  # Columns A-N
        cell = ws.cell(row=row_num, column=col_num)
        
        if cell.value is not None:
            row_has_content = True
            col_letter = get_column_letter(col_num)
            cell_ref = f"{col_letter}{row_num}"
            
            fmt = analyze_cell_format(cell)
            
            # Format value for display
            value_str = str(cell.value)
            if len(value_str) > 40:
                value_str = value_str[:37] + "..."
            
            # Build format string
            fmt_parts = []
            if fmt['font_bold']:
                fmt_parts.append('BOLD')
            if fmt['font_italic']:
                fmt_parts.append('ITALIC')
            if fmt['font_size']:
                fmt_parts.append(f"Size:{fmt['font_size']}")
            if fmt['font_color'] and fmt['font_color'] != '00000000':
                fmt_parts.append(f"Color:{fmt['font_color']}")
            if fmt['fill_color'] and fmt['fill_color'] not in ['00000000', 'None']:
                fmt_parts.append(f"Fill:{fmt['fill_color']}")
            if fmt['alignment_horizontal']:
                fmt_parts.append(f"H:{fmt['alignment_horizontal']}")
            if fmt['alignment_vertical']:
                fmt_parts.append(f"V:{fmt['alignment_vertical']}")
            if fmt['alignment_wrap']:
                fmt_parts.append('WRAP')
            
            fmt_str = ', '.join(fmt_parts) if fmt_parts else 'default'
            
            row_info.append(f"  {cell_ref}: '{value_str}'")
            row_info.append(f"       [{fmt_str}]")
    
    if row_has_content:
        print(f"Row {row_num}:")
        for info in row_info:
            print(info)
        print()

# Analyze merged cells
print("=" * 80)
print("MERGED CELLS (First 60 rows)")
print("=" * 80)
print()

merged_in_range = []
for merged_range in ws.merged_cells.ranges:
    # Check if merge is in first 60 rows
    if merged_range.min_row <= 60:
        merged_in_range.append(str(merged_range))

for merge in sorted(merged_in_range):
    print(f"  {merge}")

print()

# Column widths
print("=" * 80)
print("COLUMN WIDTHS")
print("=" * 80)
print()

for col_num in range(1, 15):
    col_letter = get_column_letter(col_num)
    width = ws.column_dimensions[col_letter].width
    print(f"  {col_letter}: {width}")

print()

# Row heights (first 60 rows)
print("=" * 80)
print("ROW HEIGHTS (Rows 1-60)")
print("=" * 80)
print()

for row_num in range(1, 61):
    height = ws.row_dimensions[row_num].height
    if height:
        print(f"  Row {row_num}: {height}")

print()
print("=" * 80)
print("ANALYSIS COMPLETE")
print("=" * 80)

wb.close()
