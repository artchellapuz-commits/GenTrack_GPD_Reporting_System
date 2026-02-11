"""
Create a clean sample Excel file with ONLY the required columns
"""
import openpyxl
from datetime import datetime, timedelta
from pathlib import Path

# Create a new workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Data"

# Add ONLY the required headers - NO title rows, NO extra columns
headers = [
    'Date',
    'Unit Number',
    'Generation kWh',
    'Operating Hours',
    'Availability Hours',
    'Forced Outage Hours',
    'Scheduled Outage Hours',
    'Remarks'
]

# Add headers to first row
for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num)
    cell.value = header
    cell.font = openpyxl.styles.Font(bold=True)

# Add sample data for AGUS1 (4 units) for 5 days
start_date = datetime(2026, 2, 1)
row_num = 2

for day in range(5):
    current_date = start_date + timedelta(days=day)
    for unit in range(1, 5):  # Units 1-4
        ws.cell(row=row_num, column=1).value = current_date
        ws.cell(row=row_num, column=2).value = unit
        ws.cell(row=row_num, column=3).value = 500000 + (unit * 10000)
        ws.cell(row=row_num, column=4).value = 22.5
        ws.cell(row=row_num, column=5).value = 23.0
        ws.cell(row=row_num, column=6).value = 0.5
        ws.cell(row=row_num, column=7).value = 0.0
        ws.cell(row=row_num, column=8).value = f'Normal operation - Day {day+1}'
        row_num += 1

# Format the date column
for row in range(2, row_num):
    ws.cell(row=row, column=1).number_format = 'YYYY-MM-DD'

# Auto-adjust column widths
for col_num in range(1, len(headers) + 1):
    column_letter = openpyxl.utils.get_column_letter(col_num)
    max_length = len(headers[col_num - 1])
    for row in range(2, row_num):
        cell_value = str(ws.cell(row=row, column=col_num).value)
        if len(cell_value) > max_length:
            max_length = len(cell_value)
    ws.column_dimensions[column_letter].width = max_length + 2

# Save the file
output_dir = Path(__file__).parent / 'sample_data'
output_dir.mkdir(exist_ok=True)
output_file = output_dir / 'AGUS1_Clean_Sample.xlsx'

wb.save(output_file)
print(f"✅ Clean sample Excel file created: {output_file}")
print(f"\n📊 File structure:")
print(f"   - Row 1: Headers ONLY")
print(f"   - Rows 2-21: Data (20 records)")
print(f"   - Columns: 8 (all required columns)")
print(f"\n📝 Contents:")
print(f"   - 5 days of data (Feb 1-5, 2026)")
print(f"   - 4 units (AGUS1 Units 1-4)")
print(f"   - Total: 20 records")
print(f"\n🎯 Upload instructions:")
print(f"   1. Go to http://localhost:8080/")
print(f"   2. Select plant: AGUS1")
print(f"   3. Upload file: {output_file.name}")
print(f"   4. Expected result: Success! 20 records imported.")
