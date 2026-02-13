"""
PSR Report Analyzer
Analyzes the structure of PSR REPORT Excel files to help with import
"""

import openpyxl
from datetime import datetime, date
import sys
import os


def analyze_excel_file(file_path):
    """Analyze the structure of an Excel file"""
    
    print("=" * 70)
    print(f"ANALYZING: {os.path.basename(file_path)}")
    print("=" * 70)
    print()
    
    try:
        workbook = openpyxl.load_workbook(file_path, data_only=True)
        
        print(f"📊 Number of sheets: {len(workbook.sheetnames)}")
        print(f"📄 Sheet names: {', '.join(workbook.sheetnames)}")
        print()
        
        for sheet_name in workbook.sheetnames:
            print(f"\n{'=' * 70}")
            print(f"SHEET: {sheet_name}")
            print('=' * 70)
            
            sheet = workbook[sheet_name]
            
            # Get dimensions
            max_row = sheet.max_row
            max_col = sheet.max_column
            print(f"📏 Dimensions: {max_row} rows × {max_col} columns")
            print()
            
            # Analyze first 20 rows
            print("🔍 First 20 rows preview:")
            print("-" * 70)
            
            for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=20), start=1):
                row_data = []
                for cell in row[:10]:  # First 10 columns
                    if cell.value is not None:
                        value = str(cell.value)
                        if len(value) > 15:
                            value = value[:12] + "..."
                        row_data.append(value)
                    else:
                        row_data.append("")
                
                if any(row_data):  # Only show non-empty rows
                    print(f"Row {row_idx:2d}: {' | '.join(row_data)}")
            
            print()
            
            # Detect format
            print("🔎 Format Detection:")
            print("-" * 70)
            
            # Look for plant codes
            plant_codes = []
            for row in sheet.iter_rows(min_row=1, max_row=50):
                first_cell = row[0].value
                if first_cell:
                    text = str(first_cell).strip().upper()
                    if len(text) <= 10 and text not in ['PLANT', 'CODE', 'TOTAL', 'SUBTOTAL']:
                        if any(c.isalpha() for c in text) and any(c.isdigit() for c in text):
                            plant_codes.append(text)
            
            if plant_codes:
                print(f"✓ Found {len(plant_codes)} potential plant codes")
                print(f"  Examples: {', '.join(plant_codes[:5])}")
            else:
                print("✗ No plant codes detected")
            
            # Look for dates
            dates_found = []
            for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=20), start=1):
                for col_idx, cell in enumerate(row, start=1):
                    if cell.value:
                        # Check if it's a date
                        if isinstance(cell.value, (datetime, date)):
                            dates_found.append((row_idx, col_idx, cell.value))
                        # Check if it's a day number
                        elif isinstance(cell.value, (int, float)) and 1 <= cell.value <= 31:
                            dates_found.append((row_idx, col_idx, f"Day {int(cell.value)}"))
            
            if dates_found:
                print(f"✓ Found {len(dates_found)} date/day columns")
                print(f"  Examples:")
                for row, col, val in dates_found[:5]:
                    print(f"    Row {row}, Col {col}: {val}")
            else:
                print("✗ No date columns detected")
            
            # Look for month/year
            month_year = None
            for row in sheet.iter_rows(min_row=1, max_row=15):
                for cell in row:
                    if cell.value:
                        text = str(cell.value).upper()
                        if any(m in text for m in ['JANUARY', 'FEBRUARY', 'MARCH', 'APRIL', 'MAY', 'JUNE',
                                                    'JULY', 'AUGUST', 'SEPTEMBER', 'OCTOBER', 'NOVEMBER', 'DECEMBER']):
                            month_year = text
                            break
                if month_year:
                    break
            
            if month_year:
                print(f"✓ Found month/year: {month_year}")
            else:
                print("✗ No month/year information detected")
            
            # Detect format type
            print()
            print("📋 Detected Format:")
            if dates_found and plant_codes:
                if any('Day' in str(d[2]) for d in dates_found):
                    print("  → DAILY FORMAT (1DATA APAO style)")
                else:
                    print("  → PSR REPORT FORMAT")
            else:
                print("  → UNKNOWN FORMAT")
            
            print()
            
            # Sample data values
            print("📊 Sample Data Values:")
            print("-" * 70)
            
            sample_values = []
            for row in sheet.iter_rows(min_row=10, max_row=30):
                for cell in row[1:10]:  # Skip first column
                    if cell.value is not None:
                        if isinstance(cell.value, (int, float)):
                            sample_values.append(cell.value)
            
            if sample_values:
                print(f"  Found {len(sample_values)} numeric values")
                print(f"  Range: {min(sample_values):.2f} to {max(sample_values):.2f}")
                print(f"  Average: {sum(sample_values) / len(sample_values):.2f}")
                print(f"  Sample values: {', '.join(str(v) for v in sample_values[:10])}")
            else:
                print("  No numeric values found")
            
            print()
    
    except Exception as e:
        print(f"❌ ERROR: {str(e)}")
        import traceback
        traceback.print_exc()
    
    print()
    print("=" * 70)
    print("ANALYSIS COMPLETE")
    print("=" * 70)


def main():
    """Main function"""
    
    # Default file path
    default_file = r"REPORTS\8. PSR\PSR REPORT-8AM.xlsx"
    
    if len(sys.argv) > 1:
        file_path = sys.argv[1]
    else:
        file_path = default_file
    
    if not os.path.exists(file_path):
        print(f"❌ File not found: {file_path}")
        print()
        print("Usage:")
        print(f"  python {sys.argv[0]} <path_to_excel_file>")
        print()
        print("Example:")
        print(f'  python {sys.argv[0]} "REPORTS\\8. PSR\\PSR REPORT-8AM.xlsx"')
        return
    
    analyze_excel_file(file_path)


if __name__ == "__main__":
    main()
